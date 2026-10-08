#!/usr/bin/env python3
"""
Genera nóminas de prueba completando la plantilla real "Formato envío notas
Liceos" con alumnos inventados. Sirve para probar la carga sin usar datos
reales de estudiantes.

Uso:
    python3 generar_nominas_demo.py <plantilla.xlsx> <carpeta_salida>

Genera un archivo por cada liceo de prueba. El liceo 05 trae errores a
propósito (nota con coma, nota sin coma, alumno sin notas, duplicado) para
ver cómo se muestran en la previsualización.
"""
import random
import sys
from pathlib import Path

import openpyxl

NOMBRES = ["Sofía", "Martina", "Florencia", "Isidora", "Agustina", "Josefa", "Emilia", "Catalina",
           "Antonia", "Valentina", "Mateo", "Benjamín", "Vicente", "Tomás", "Agustín", "Joaquín",
           "Maximiliano", "Lucas", "Matías", "Cristóbal", "Fernanda", "Javiera", "Camila", "Diego"]
APELLIDOS = ["González", "Muñoz", "Rojas", "Díaz", "Pérez", "Soto", "Contreras", "Silva", "Martínez",
             "Sepúlveda", "Morales", "Rodríguez", "López", "Fuentes", "Hernández", "Torres", "Araya",
             "Flores", "Espinoza", "Valenzuela", "Castillo", "Tapia", "Reyes", "Gutiérrez", "Castro"]

LICEOS_DEMO = [
    ("01", "Colegio Parroquial de Andacollo", ["3° año A", "3° año B"]),
    ("02", "Liceo Pedro Regalado Videla", ["3° año A", "3° año B", "3° C"]),
    ("04", "Colegio Pablo Neruda", ["3°A", "3°B"]),
    ("05", "Instituto de Ad. y Com. Estado de Israel IAC", ["3° A", "3° B", "3° C"]),
]


def nota(rnd):
    return round(min(7.0, max(4.0, rnd.gauss(5.8, 0.55))), 1)


def main(plantilla, salida):
    salida = Path(salida)
    salida.mkdir(parents=True, exist_ok=True)
    rnd = random.Random(2026)
    for codigo, liceo, cursos in LICEOS_DEMO:
        wb = openpyxl.load_workbook(plantilla)
        hojas = wb.worksheets
        for i, ws in enumerate(hojas):
            if i >= len(cursos):
                wb.remove(ws)
                continue
            ws.title = f"Curso {cursos[i].replace(' año ', '').replace(' ', '')}"
            ws["C2"] = liceo
            ws["C3"] = "Contraparte de prueba"
            ws["C4"] = "Orientadora"
            ws["C6"] = cursos[i]
            ws["C7"] = "Científico Humanista" if i == 0 else "Técnico Profesional"
            n_alumnos = rnd.randint(22, 34)
            for k in range(n_alumnos):
                fila = 9 + k
                ws.cell(fila, 1, k + 1)
                ws.cell(fila, 2, rnd.choice(APELLIDOS))
                ws.cell(fila, 3, rnd.choice(APELLIDOS))
                ws.cell(fila, 4, f"{rnd.choice(NOMBRES)} {rnd.choice(NOMBRES)}")
                ws.cell(fila, 5, nota(rnd))
                ws.cell(fila, 6, nota(rnd))
            if codigo == "05" and i == 0:
                # Errores típicos que comete un liceo al llenar la planilla
                ws.cell(9, 5, "6,8")                       # coma decimal en texto
                ws.cell(10, 6, 65)                         # nota sin coma
                # (ws.cell(f, c, None) no borra en openpyxl: hay que asignar .value)
                ws.cell(11, 5).value = None                # alumno sin notas
                ws.cell(11, 6).value = None
                ws.cell(12, 6).value = None                # falta un promedio
                for col in range(2, 7):                    # alumno duplicado
                    ws.cell(14, col, ws.cell(13, col).value)
                ws.cell(15, 5, 8.2)                        # nota fuera de rango
        destino = salida / f"nomina_liceo_{codigo}.xlsx"
        wb.save(destino)
        print("generado", destino)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
